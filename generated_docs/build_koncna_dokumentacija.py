from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "generated_docs"
DOCX_PATH = OUT_DIR / "Koncna_dokumentacija_VozimVarno.docx"
DIAGRAM_DIR = OUT_DIR / "koncna_diagrami"

SCREENSHOTS = [
    ("Domov", "Glavni meni aplikacije s hitrim dostopom do nove vožnje, zgodovine, rezultatov, ocen, izzivov in nastavitev.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073044.png")),
    ("Nova vožnja - dovoljenja", "Pred začetkom vožnje aplikacija preveri kamero, GPS, pospeškometer, mikrofon in shranjevanje.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073051.png")),
    ("Vožnja v teku", "Zaslon med vožnjo prikazuje čas, hitrost, dogodke pospeševanja/zaviranja/odstopanja hitrosti ter stanje GPS in hrupa.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073105.png")),
    ("Povzetek vožnje", "Po zaključku se prikaže varnostna ocena, razdalja, čas vožnje, povprečna hitrost in število dogodkov.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073113.png")),
    ("Moje vožnje", "Zgodovina omogoča pregled preteklih voženj, filtriranje po obdobjih in hitro primerjavo ocen.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073206.png")),
    ("Profil", "Profil prikazuje osnovne podatke voznika, statistiko, dosežke, nastavitve, pomoč in odjavo.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073214.png")),
    ("Rezultati", "Analitični zaslon prikazuje povprečno varnostno oceno, število voženj, razdaljo, dogodke in trende.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073222.png")),
    ("Ocena vožnje", "Uporabnik lahko vožnjo oceni z zvezdicami in kasneje odpre podrobnosti posamezne vožnje.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073244.png")),
    ("Začetni zaslon", "Splash zaslon predstavi blagovno znamko VozimVarno in slogan projekta.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073254.png")),
    ("Nastavitve", "Nastavitve omogočajo urejanje računa, obvestil, enot, teme, dovoljenj in pragov incidentov.",
     Path(r"C:/Users/Filip/Pictures/Screenshots/Screenshot 2026-06-04 073329.png")),
]


def font(size: int, bold: bool = False):
    names = ["arialbd.ttf" if bold else "arial.ttf", "calibrib.ttf" if bold else "calibri.ttf"]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def rounded(draw, box, fill, outline="#D8E1EC", width=2, radius=18):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def centered_text(draw, box, text, fnt, fill="#0B2545"):
    lines = text.split("\n")
    line_heights = [draw.textbbox((0, 0), line, font=fnt)[3] for line in lines]
    total = sum(line_heights) + (len(lines) - 1) * 8
    y = box[1] + ((box[3] - box[1]) - total) / 2
    for line, h in zip(lines, line_heights):
        w = draw.textbbox((0, 0), line, font=fnt)[2]
        draw.text((box[0] + (box[2] - box[0] - w) / 2, y), line, font=fnt, fill=fill)
        y += h + 8


def arrow(draw, start, end, color="#1F6FB8", width=4):
    draw.line([start, end], fill=color, width=width)
    import math
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 13
    pts = [
        end,
        (end[0] - size * math.cos(ang - 0.45), end[1] - size * math.sin(ang - 0.45)),
        (end[0] - size * math.cos(ang + 0.45), end[1] - size * math.sin(ang + 0.45)),
    ]
    draw.polygon(pts, fill=color)


def make_use_case_diagram(path: Path):
    img = Image.new("RGB", (1400, 900), "#FFFFFF")
    d = ImageDraw.Draw(img)
    title = font(38, True)
    h = font(26, True)
    t = font(22)
    d.text((40, 35), "Diagram primerov uporabe - VozimVarno", font=title, fill="#0B2545")
    rounded(d, (435, 105, 1320, 820), "#F8FAFC", "#BFD0E2", 3, 28)
    d.text((480, 125), "Mobilna aplikacija", font=h, fill="#0B2545")
    actor = (120, 420)
    d.ellipse((actor[0]-35, actor[1]-120, actor[0]+35, actor[1]-50), outline="#0B2545", width=5)
    d.line((actor[0], actor[1]-50, actor[0], actor[1]+95), fill="#0B2545", width=5)
    d.line((actor[0]-70, actor[1], actor[0]+70, actor[1]), fill="#0B2545", width=5)
    d.line((actor[0], actor[1]+95, actor[0]-60, actor[1]+190), fill="#0B2545", width=5)
    d.line((actor[0], actor[1]+95, actor[0]+60, actor[1]+190), fill="#0B2545", width=5)
    centered_text(d, (40, 635, 240, 700), "Voznik", h)
    use_cases = [
        (620, 210, "Ustvari račun\nin prijava"),
        (980, 220, "Nastavi dovoljenja\nin pragove"),
        (650, 400, "Začne in spremlja\nvožnjo"),
        (1030, 425, "Zaključi in shrani\nvožnjo"),
        (650, 610, "Pregleda zgodovino\nin podrobnosti"),
        (1030, 640, "Oceni vožnjo\nin spremlja rezultate"),
    ]
    for x, y, label in use_cases:
        d.ellipse((x-155, y-70, x+155, y+70), fill="#FFFFFF", outline="#1F6FB8", width=4)
        centered_text(d, (x-140, y-50, x+140, y+50), label, t)
        arrow(d, (245, 485), (x-155, y), "#6B7A90", 3)
    d.text((445, 830), "Glavni uporabnik je voznik, ki upravlja vožnjo, dovoljenja, zgodovino, ocene in nastavitve.", font=t, fill="#4A5A70")
    img.save(path)


def make_architecture_diagram(path: Path):
    img = Image.new("RGB", (1400, 900), "#FFFFFF")
    d = ImageDraw.Draw(img)
    title = font(38, True)
    h = font(24, True)
    small = font(20)
    d.text((40, 35), "Diagram arhitekture sistema", font=title, fill="#0B2545")
    boxes = {
        "ui": (70, 155, 390, 300, "React Native / Expo UI\nzasloni in navigacija"),
        "session": (510, 145, 880, 315, "useRideSession\nlogika aktivne vožnje"),
        "sensors": (70, 430, 390, 610, "Naprava\nGPS, kamera, mikrofon,\npospeškometer, žiroskop"),
        "local": (510, 430, 880, 610, "Lokalna shramba\nAsyncStorage,\nnastavitve in vožnje"),
        "convex": (1010, 145, 1330, 315, "Convex backend\nAuth, queries, mutations"),
        "db": (1010, 430, 1330, 610, "Podatkovna baza\nusers, rides,\nridePoints, incidents"),
        "weather": (510, 705, 880, 815, "Open-Meteo API\nvremenski kontekst"),
    }
    for key, (x1, y1, x2, y2, txt) in boxes.items():
        fill = "#F4F7FB" if key not in ["convex", "db"] else "#EEF7F1"
        rounded(d, (x1, y1, x2, y2), fill, "#BFD0E2", 3, 20)
        centered_text(d, (x1+12, y1+10, x2-12, y2-10), txt, h if key in ["ui", "session", "convex", "db"] else small)
    arrow(d, (390, 225), (510, 225))
    arrow(d, (230, 430), (230, 300))
    arrow(d, (390, 520), (510, 520))
    arrow(d, (880, 225), (1010, 225))
    arrow(d, (1170, 315), (1170, 430))
    arrow(d, (690, 610), (690, 705))
    arrow(d, (880, 520), (1010, 520))
    d.text((60, 835), "Pretok: uporabniški vmesnik vodi sejo vožnje, senzorji ustvarjajo meritve, lokalna shramba omogoča offline delovanje, Convex sinhronizira podatke uporabnika.", font=small, fill="#4A5A70")
    img.save(path)


def make_er_diagram(path: Path):
    img = Image.new("RGB", (1400, 930), "#FFFFFF")
    d = ImageDraw.Draw(img)
    title = font(38, True)
    h = font(24, True)
    t = font(18)
    d.text((40, 35), "Podatkovni model / E-R diagram", font=title, fill="#0B2545")
    entities = {
        "users": (80, 150, 360, 295, ["_id", "email/password auth", "ime profila"]),
        "userSettings": (80, 480, 405, 705, ["userId", "notificationsEnabled", "speedUnit, theme", "camera/gps/microphone", "incident thresholds", "updatedAt"]),
        "rides": (560, 125, 930, 375, ["userId", "clientRideId", "startTime, endTime", "durationSeconds", "distanceKm", "score", "max/avgSpeedKmh", "rating/comment", "weather fields"]),
        "ridePoints": (1060, 110, 1340, 325, ["rideId", "userId", "latitude", "longitude", "speedKmh", "timestamp", "altitude?"]),
        "rideIncidents": (1060, 500, 1340, 755, ["rideId", "userId", "type", "timestamp", "intensity", "speedKmh?", "lat/lon?"]),
    }
    for name, (x1, y1, x2, y2, fields) in entities.items():
        rounded(d, (x1, y1, x2, y2), "#F8FAFC", "#1F6FB8", 3, 18)
        d.rectangle((x1, y1, x2, y1+48), fill="#0B4F8F")
        d.text((x1+16, y1+13), name, font=h, fill="#FFFFFF")
        y = y1 + 62
        for fld in fields:
            d.text((x1+18, y), fld, font=t, fill="#0B2545")
            y += 27
    arrow(d, (360, 215), (560, 230), "#6B7A90", 4)
    d.text((425, 190), "1:N", font=t, fill="#6B7A90")
    arrow(d, (245, 295), (245, 480), "#6B7A90", 4)
    d.text((260, 385), "1:1", font=t, fill="#6B7A90")
    arrow(d, (930, 235), (1060, 220), "#6B7A90", 4)
    d.text((970, 190), "1:N", font=t, fill="#6B7A90")
    arrow(d, (930, 330), (1060, 610), "#6B7A90", 4)
    d.text((975, 490), "1:N", font=t, fill="#6B7A90")
    d.text((60, 850), "Indeksi v Convex shemi: userId, userId_clientRideId, userId_startTime in rideId. Ti indeksi pospešijo zgodovino, sinhronizacijo in prikaz podrobnosti.", font=t, fill="#4A5A70")
    img.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_doc(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        s = styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    run = p.add_run("VozimVarno")
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = RGBColor.from_string("0B4F8F")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("Končna dokumentacija projekta").bold = True
    p2.runs[0].font.size = Pt(18)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("React Native / Expo mobilna aplikacija za varnejšo vožnjo")
    p3.runs[0].font.size = Pt(13)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(40)
    p4.add_run("Izbrane kategorije: primeri uporabe, podatkovna baza, zaslonske slike, arhitektura in varnost.")
    doc.add_page_break()


def add_bullets(doc: Document, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        hdr[i].width = Inches(widths[i])
        set_cell_shading(hdr[i], "F2F4F7")
        set_cell_margins(hdr[i])
        for p in hdr[i].paragraphs:
            p.runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
    return table


def add_screenshot_grid(doc: Document):
    doc.add_heading("3. Zaslonske slike končnega grafičnega vmesnika", level=1)
    doc.add_paragraph(
        "Končni uporabniški vmesnik je zasnovan kot mobilna aplikacija z jasnimi karticami, spodnjo navigacijo in namenskimi zasloni za vožnjo, zgodovino, rezultate, ocene, profil in nastavitve."
    )
    valid = [(title, desc, path) for title, desc, path in SCREENSHOTS if path.exists()]
    if len(valid) < len(SCREENSHOTS):
        missing = [str(path) for _, _, path in SCREENSHOTS if not path.exists()]
        doc.add_paragraph("Opomba: manjkajoče slike niso bile vstavljene: " + "; ".join(missing))

    for idx, (title, desc, path) in enumerate(valid, 1):
        if idx > 1 and (idx - 1) % 2 == 0:
            doc.add_page_break()
        doc.add_heading(f"Slika {idx}: {title}", level=2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width = Inches(2.35)
        if path.name.endswith("_details.png"):
            width = Inches(2.7)
        p.add_run().add_picture(str(path), width=width)
        cap = doc.add_paragraph(desc)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)


def build_doc():
    OUT_DIR.mkdir(exist_ok=True)
    DIAGRAM_DIR.mkdir(exist_ok=True)
    use_case = DIAGRAM_DIR / "diagram_primerov_uporabe.png"
    arch = DIAGRAM_DIR / "diagram_arhitekture.png"
    er = DIAGRAM_DIR / "er_diagram.png"
    make_use_case_diagram(use_case)
    make_architecture_diagram(arch)
    make_er_diagram(er)

    doc = Document()
    style_doc(doc)
    add_title(doc)

    doc.add_heading("1. Diagram primerov uporabe", level=1)
    doc.add_paragraph(
        "Aplikacija VozimVarno je namenjena vozniku, ki z mobilno napravo spremlja varnost vožnje, pregleduje rezultate in prilagaja pragove zaznavanja incidentov."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(use_case), width=Inches(6.3))
    add_bullets(doc, [
        "Voznik pred vožnjo preveri dovoljenja za kamero, GPS, pospeškometer, mikrofon in shranjevanje.",
        "Med vožnjo sistem beleži hitrost, lokacijske točke in incidente, kot so močno pospeševanje, zaviranje, ostri zavoji, prekoračitev hitrosti in hrup.",
        "Po vožnji voznik pregleda povzetek, shrani vožnjo, oceni izkušnjo in spremlja trende varnosti.",
    ])

    doc.add_heading("2. Podatkovna baza", level=1)
    doc.add_paragraph(
        "Podatkovni model temelji na Convex shemi. Glavna entiteta je vožnja, ki pripada uporabniku in ima povezane GPS točke ter dogodke. Nastavitve uporabnika so shranjene ločeno, ker se uporabljajo za dovoljenja, temo, enote in pragove incidentov."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(er), width=Inches(6.3))
    add_table(doc, ["Tabela", "Namen", "Ključna polja"], [
        ("users", "Avtenticirani uporabniki sistema.", "id, podatki Convex Auth"),
        ("userSettings", "Shrani uporabniške nastavitve in pragove incidentov.", "userId, speedUnit, theme, cameraEnabled, gpsEnabled, microphoneEnabled, thresholds"),
        ("rides", "Osnovni zapis vožnje in povzetek rezultatov.", "userId, startTime, durationSeconds, distanceKm, score, avgSpeedKmh, userRating, weather"),
        ("ridePoints", "GPS pot vožnje po časovnih točkah.", "rideId, latitude, longitude, speedKmh, timestamp"),
        ("rideIncidents", "Zaznani dogodki med vožnjo.", "rideId, type, timestamp, intensity, speedKmh, latitude, longitude"),
    ], [1.25, 2.25, 3.0])

    add_screenshot_grid(doc)

    doc.add_page_break()
    doc.add_heading("4. Diagram arhitekture sistema", level=1)
    doc.add_paragraph(
        "Arhitektura je razdeljena na mobilni odjemalec, lokalno shrambo, zunanji vremenski API in Convex backend. Takšna zgradba omogoča prikaz podatkov tudi ob šibkejši povezavi, nato pa sinhronizacijo z uporabniškim računom."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(arch), width=Inches(6.3))
    add_table(doc, ["Komponenta", "Vloga"], [
        ("React Native / Expo", "Prikazuje zaslone, upravlja navigacijo in bere dovoljenja naprave."),
        ("useRideSession", "Koordinira aktivno vožnjo, časovnik, senzorje, dogodke, rezultat in shranjevanje."),
        ("Senzorske storitve", "GPS, pospeškometer, žiroskop, kamera in mikrofon posredujejo meritve za izračun incidentov."),
        ("AsyncStorage", "Hrani lokalne vožnje, nastavitve in vremenski predpomnilnik."),
        ("Convex backend", "Izvaja avtentikacijo, preverjanje lastništva podatkov, mutacije, poizvedbe in trajno hrambo."),
        ("Open-Meteo", "Doda vremenski kontekst ob koncu vožnje."),
    ], [2.0, 4.5])

    doc.add_heading("5. Varnostne znacilnosti", level=1)
    doc.add_paragraph(
        "Varnost je obravnavana na ravni prijave, dovoljenj naprave, dostopa do podatkov in odpornosti aplikacije pri lokalnem shranjevanju."
    )
    add_bullets(doc, [
        "Avtentikacija: projekt uporablja Convex Auth z geselnim ponudnikom. Backend za mutacije in poizvedbe pridobi `userId` iz seje.",
        "Avtorizacija podatkov: funkcija `requireRideOwner` preveri, da uporabnik dostopa samo do svojih voženj; pri napaki vrne sporočilo, kot da vožnja ne obstaja.",
        "Dovoljenja naprave: zaslon pred vožnjo preveri kamero, GPS, mikrofon, pospeškometer in shranjevanje. Uporabnik mora manjkajoča dovoljenja odobriti pred začetkom.",
        "Omejitev obdelave: kamera in mikrofon se aktivirata le, če sta omogočena v nastavitvah. Ob zaključku vožnje se senzorji, sledenje lokaciji in monitoring hrupa ustavijo.",
        "Zaščita pred podvajanjem: sinhronizacija uporablja `clientRideId` in indeks `userId_clientRideId`, zato se lokalna vožnja pri ponovni sinhronizaciji posodobi namesto podvoji.",
        "Lokalna odpornost: če backend ali vremenski API ni dosegljiv, aplikacija ohrani lokalno vožnjo kot vir resnice in uporabi predpomnilnik ali fallback za vreme.",
    ])
    doc.add_paragraph(
        "Priporočilo za produkcijo: obdelavo občutljivih lokacijskih podatkov je smiselno dopolniti z jasnim pravilnikom zasebnosti, krajšim rokom hrambe poti in dodatnim šifriranjem lokalnih podatkov, kjer platforma to omogoča."
    )

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("VozimVarno - končna dokumentacija")

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
    print(DOCX_PATH)
